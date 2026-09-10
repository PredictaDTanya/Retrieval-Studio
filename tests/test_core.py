"""v2 test suite (stdlib unittest — run with:  python -m unittest discover tests -v)

Covers the spec §14 list: extraction shapes, dataset isolation, model
integrity, entity matching, prompt validation, reconciliation, raw
persistence + reprocess, unbranded studies, failed responses.
All data comes from tests/fixtures.py (labelled TEST FIXTURES).
"""

from __future__ import annotations

import io
import json
import os
import shutil
import tempfile
import time
import unittest

import core.datasets as ds
from core import runner, analysis, exports
from core.entities import EntityMatcher, new_entity
from core.extraction import extract_observation
from core.classify import classify_domain, classify_query, TAXONOMY
from core.model import canonical_url, open_db
from core.prompts import from_lines, parse_upload, validate, new_prompt
from tests import fixtures as fx


def _matcher():
    return EntityMatcher(fx.SAMPLE_ENTITIES())


class TestExtraction(unittest.TestCase):
    def setUp(self):
        self.ex = extract_observation(fx.rich_payload(), _matcher())

    def test_queries_plural_singular_missing(self):
        qs = self.ex["fanout_queries"]
        texts = [q["query"] for q in qs if not q["query_missing"]]
        self.assertIn("best SEO agency Springfield", texts)          # queries[]
        self.assertIn("Springfield digital marketing pricing", texts)  # query
        missing = [q for q in qs if q["query_missing"]]
        self.assertEqual(len(missing), 1)          # search w/o query text
        self.assertIsNone(missing[0]["query"])     # never invented

    def test_site_operator(self):
        site = [q for q in self.ex["fanout_queries"] if q["has_site_operator"]]
        self.assertEqual(len(site), 1)
        self.assertEqual(site[0]["site_domain"], "acmeseo.example")

    def test_sources_opened_inpage_citations_distinct(self):
        self.assertEqual(len(self.ex["search_sources"]), 4)  # incl. feed
        types = {s["source_type"] for s in self.ex["search_sources"]}
        self.assertIn("oai-weather", types)   # typed feed source preserved
        self.assertIn("url", types)
        self.assertEqual(len(self.ex["opened_pages"]), 1)
        self.assertEqual(len(self.ex["inpage_searches"]), 1)
        self.assertEqual(self.ex["inpage_searches"][0]["pattern"], "pricing")
        # citations: 3 annotations, 2 unique canonical urls
        self.assertEqual(self.ex["rollups"]["num_citations"], 3)
        self.assertEqual(self.ex["rollups"]["num_unique_cited_urls"], 2)

    def test_url_canonicalisation_preserves_raw(self):
        acme = [c for c in self.ex["citations"]
                if "acmeseo" in c["canonical_url"]][0]
        self.assertNotIn("utm_source", acme["canonical_url"])
        self.assertIn("utm_source=openai", acme["raw_url"])
        self.assertEqual(acme["citation_count"], 2)  # dupe counted, not lost

    def test_reasoning_summary_only_when_returned(self):
        self.assertIn("TEST FIXTURE reasoning", self.ex["reasoning_summary"])
        plain = extract_observation(fx.plain_payload(), _matcher())
        self.assertIsNone(plain["reasoning_summary"])  # never fabricated

    def test_actual_model_captured(self):
        self.assertEqual(self.ex["actual_model"], "gpt-5.6-sol")

    def test_entity_attribution_per_stage(self):
        er = {e["entity_id"]: e for e in self.ex["entity_runs"]}
        # Acme: returned in search, page opened, cited — but NOT in-page
        # searched. Each stage is a separate fact.
        self.assertEqual(er["ent_acme"]["source_returned"], 1)
        self.assertEqual(er["ent_acme"]["page_opened"], 1)
        self.assertEqual(er["ent_acme"]["page_searched"], 0)
        self.assertEqual(er["ent_acme"]["cited"], 1)
        self.assertEqual(er["ent_acme"]["search_targeted"], 1)
        # Rival: returned, in-page searched, cited — but no page_opened.
        self.assertEqual(er["ent_rival"]["source_returned"], 1)
        self.assertEqual(er["ent_rival"]["page_opened"], 0)
        self.assertEqual(er["ent_rival"]["page_searched"], 1)
        self.assertEqual(er["ent_rival"]["cited"], 1)
        self.assertEqual(er["ent_rival"]["named_in_answer"], 1)
        self.assertEqual(er["ent_neutral"]["mentioned_anywhere"], 0)
        rels = {(r["entity_id"], r["relation"])
                for r in self.ex["source_entity"]}
        self.assertIn(("ent_acme", "first_party"), rels)

    def test_response_status_captured(self):
        self.assertEqual(self.ex["response_status"], "completed")
        inc = extract_observation(fx.incomplete_payload(), _matcher())
        self.assertEqual(inc["response_status"], "incomplete")
        self.assertEqual(inc["incomplete_reason"], "max_output_tokens")

    def test_no_entity_unbranded(self):
        ex = extract_observation(fx.rich_payload(), None)
        self.assertEqual(ex["entity_runs"], [])
        self.assertEqual(ex["source_entity"], [])
        self.assertEqual(ex["rollups"]["num_citations"], 3)


class TestEntities(unittest.TestCase):
    def test_alias_and_domain_matching(self):
        m = _matcher()
        self.assertEqual(m.match_text("I would pick Acme for this"),
                         ["ent_acme"])
        self.assertEqual(m.match_domain("blog.acmeseo.example"), ["ent_acme"])
        self.assertEqual(m.match_text("nothing relevant"), [])

    def test_no_auto_competitor(self):
        e = new_entity("Someone", role="not-a-role")
        self.assertEqual(e["role"], "comparison")  # safe default, never competitor


class TestClassify(unittest.TestCase):
    def test_domain_classes(self):
        self.assertEqual(classify_domain("clutch.co"), "directory")
        self.assertEqual(classify_domain("trustpilot.com"), "review_platform")
        self.assertEqual(classify_domain("consumer.vic.gov.au"), "government")
        self.assertEqual(classify_domain("acmeseo.example",
                                         ["acmeseo.example"]), "first_party")
        self.assertEqual(classify_domain("random-blog.example"),
                         "other_third_party")

    def test_taxonomy_stable(self):
        self.assertEqual(classify_query("best SEO agencies in Melbourne 2026"),
                         "Agency discovery")
        self.assertEqual(classify_query("Acme reviews", _matcher()),
                         "Direct entity investigation")
        self.assertEqual(
            classify_query("site:acmeseo.example services", _matcher(),
                           "acmeseo.example"),
            "Official / first-party verification")
        for q in ("chatgpt seo visibility", "shopify seo experts",
                  "seo pricing melbourne"):
            self.assertIn(classify_query(q), TAXONOMY)


class TestLocationParsing(unittest.TestCase):
    def test_flexible_forms(self):
        f = ds.parse_location_string
        self.assertEqual(f("Melbourne, Victoria, AU"),
                         {"city": "Melbourne", "region": "Victoria",
                          "country": "AU"})
        self.assertEqual(f("Sydney, AU"),
                         {"city": "Sydney", "country": "AU"})
        self.assertEqual(f("Australia"), {"country": "AU"})   # country name
        self.assertEqual(f("au"), {"country": "AU"})          # code, cased
        self.assertEqual(f("United Kingdom"), {"country": "GB"})
        self.assertIsNone(f(""))
        self.assertIsNone(f("   "))

    def test_format_roundtrip(self):
        self.assertEqual(ds.format_location({"country": "AU"}), "AU")
        self.assertEqual(ds.format_location(
            {"city": "Sydney", "country": "AU"}), "Sydney, AU")
        self.assertEqual(ds.format_location(
            {"city": "Melbourne", "region": "Victoria", "country": "AU"}),
            "Melbourne, Victoria, AU")


class TestPrompts(unittest.TestCase):
    def test_list_markers_stripped(self):
        ps = from_lines("1. first prompt\n2) second prompt\n- third prompt\n")
        self.assertEqual([p["prompt_text"] for p in ps],
                         ["first prompt", "second prompt", "third prompt"])

    def test_bulk_and_duplicates(self):
        ps = from_lines("alpha\nbeta\nalpha\n")
        v = validate(ps)
        self.assertTrue(any("duplicates" in d for d in v["duplicate_prompts"]))
        self.assertTrue(v["ok"])  # duplicate text warns, doesn't block

    def test_upload_validation(self):
        class F(io.BytesIO):
            name = "p.csv"
        f = F(("prompt_id,prompt_text,location,bogus\n"
               "P01,first prompt,\"Melbourne, Victoria, AU\",x\n"
               "P01,,badloc,y\n").encode("utf-8"))
        ps, unsupported, blanks = parse_upload(f)
        v = validate(ps, unsupported, blanks)
        self.assertEqual(unsupported, ["bogus"])
        self.assertEqual(blanks, 1)
        self.assertFalse(v["duplicate_ids"])  # blank row dropped before dupe


class _DataTestBase(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="fanout_v2_test_")
        self._old_root = ds.DATA_ROOT
        ds.DATA_ROOT = os.path.join(self._tmp, "studies")

    def tearDown(self):
        ds.DATA_ROOT = self._old_root
        shutil.rmtree(self._tmp, ignore_errors=True)

    def _mk_study(self, entities=True):
        return ds.create_study(
            "Test study", mode="multi" if entities else "unbranded",
            entities=fx.SAMPLE_ENTITIES() if entities else [],
            prompts=fx.SAMPLE_PROMPTS())

    def _fill(self, study, dtype="live"):
        man = ds.create_dataset(study, dtype, "gpt-5.6-sol", 2)
        m = EntityMatcher(study["entities"]) if study["entities"] else None
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        try:
            for i, (meta, raw) in enumerate(fx.sample_payloads(), 1):
                meta = dict(meta)
                meta["observation_id"] = f"obs_t{i:03d}"
                meta["requested_model"] = "gpt-5.6-sol"
                ds.store_observation(man, meta, raw, m, {}, conn=conn)
            conn.commit()
        finally:
            conn.close()
        return ds.update_totals(man)


class TestDatasetIsolation(_DataTestBase):
    def test_live_and_sample_are_separate_files(self):
        study = self._mk_study()
        live = self._fill(study, "live")
        sample = self._fill(study, "sample")
        self.assertNotEqual(live["dataset_id"], sample["dataset_id"])
        p1 = ds.dataset_db_path(study["study_id"], live["dataset_id"])
        p2 = ds.dataset_db_path(study["study_id"], sample["dataset_id"])
        self.assertNotEqual(p1, p2)
        for p, did in ((p1, live["dataset_id"]), (p2, sample["dataset_id"])):
            conn = open_db(p)
            ids = [r[0] for r in conn.execute(
                "SELECT DISTINCT dataset_id FROM observations")]
            conn.close()
            self.assertEqual(ids, [did])  # zero cross-contamination

    def test_raw_persisted_and_reprocess(self):
        study = self._mk_study()
        man = self._fill(study)
        raw_dir = os.path.join(ds.dataset_dir(study["study_id"],
                                              man["dataset_id"]), "raw")
        raws = os.listdir(raw_dir)
        self.assertEqual(len(raws), 4)  # 4 ok runs have raw; failed has none
        env = json.load(open(os.path.join(raw_dir, raws[0]),
                             encoding="utf-8"))
        self.assertIn("response", env)
        self.assertIn("_meta", env)
        n = ds.reprocess(study["study_id"], man["dataset_id"], {})
        self.assertEqual(n, 4)
        man2 = ds.load_manifest(study["study_id"], man["dataset_id"])
        self.assertEqual(man2["totals"]["observations"], 5)  # failed kept
        self.assertEqual(man2["totals"]["failed"], 1)

    def test_failed_run_preserved(self):
        study = self._mk_study()
        man = self._fill(study)
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        err = conn.execute("SELECT error FROM observations WHERE "
                           "status='error'").fetchone()[0]
        conn.close()
        self.assertIn("RateLimitError", err)


class TestModelIntegrity(_DataTestBase):
    class _FakeResp:
        def __init__(self, payload):
            self._p = payload

        def model_dump(self, mode="json"):
            return self._p

    class _FakeClient:
        """Returns a mismatched model on the 2nd call."""
        def __init__(self, payloads):
            self._payloads = payloads
            self.calls = 0
            outer = self

            class R:
                def create(self, **kw):
                    outer.calls += 1
                    p = outer._payloads[min(outer.calls - 1,
                                            len(outer._payloads) - 1)]
                    if isinstance(p, Exception):
                        raise p
                    return TestModelIntegrity._FakeResp(p)
            self.responses = R()

    def test_model_matches(self):
        self.assertTrue(runner.model_matches("gpt-5.6-sol", "gpt-5.6-sol"))
        self.assertTrue(runner.model_matches("gpt-5.6-sol",
                                             "gpt-5.6-sol-2026-07-01"))
        self.assertFalse(runner.model_matches("gpt-5.6-sol", "gpt-5.6-terra"))

    def test_no_silent_fallback_aborts_batch(self):
        study = self._mk_study(entities=False)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 2)
        client = self._FakeClient([fx.plain_payload(model="gpt-5.6-sol"),
                                   fx.mismatch_payload()])
        res = runner.run_batch(client, study, man, {})
        self.assertTrue(res["aborted"])
        self.assertIn("model_mismatch", res["abort_reason"])
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        rows = conn.execute("SELECT status, requested_model, actual_model "
                            "FROM observations ORDER BY rowid").fetchall()
        conn.close()
        self.assertEqual(rows[0]["status"], "ok")
        self.assertEqual(rows[1]["status"], "error")
        self.assertEqual(rows[1]["actual_model"], "gpt-5.6-terra")
        # after abort, no further calls under the wrong model
        self.assertEqual(client.calls, 2)

    def test_model_rejection_aborts(self):
        study = self._mk_study(entities=False)
        man = ds.create_dataset(study, "live", "gpt-9-nonexistent", 1)
        client = self._FakeClient([
            RuntimeError("Error code: 404 - model 'gpt-9-nonexistent' "
                         "does not exist")])
        res = runner.run_batch(client, study, man, {})
        self.assertTrue(res["aborted"])
        self.assertIn("rejected", res["abort_reason"])


class TestExports(_DataTestBase):
    def test_reconciliation_passes_and_workbook_builds(self):
        study = self._mk_study()
        man = self._fill(study)
        checks = exports.reconcile(study["study_id"], man["dataset_id"])
        bad = [c for c in checks if not c[1]]
        self.assertEqual(bad, [], f"failed checks: {bad}")
        path = os.path.join(self._tmp, "wb.xlsx")
        exports.build_workbook(study["study_id"], man["dataset_id"], path)
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        for sheet in exports.WORKBOOK_SHEETS:
            self.assertIn(sheet[:31], wb.sheetnames)

    def test_foreign_dataset_row_blocks_export(self):
        study = self._mk_study()
        man = self._fill(study)
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        conn.execute("UPDATE fanout_queries SET dataset_id='intruder' "
                     "WHERE rowid IN (SELECT rowid FROM fanout_queries "
                     "LIMIT 1)")
        conn.commit()
        conn.close()
        checks = exports.reconcile(study["study_id"], man["dataset_id"])
        self.assertFalse(exports.reconciliation_ok(checks))
        with self.assertRaises(RuntimeError):
            exports.build_workbook(study["study_id"], man["dataset_id"],
                                   os.path.join(self._tmp, "bad.xlsx"))

    def test_unbranded_study_analysis(self):
        study = self._mk_study(entities=False)
        man = self._fill(study)
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        ov = analysis.overview(conn)
        self.assertEqual(ov["entities_tracked"], 0)
        self.assertIsNone(ov["entity_mention_rate"])  # no fake client metrics
        self.assertEqual(analysis.entity_visibility(conn, []), [])
        self.assertEqual(analysis.citation_gaps(conn, []), [])
        conn.close()

    def test_canonical_url(self):
        self.assertEqual(
            canonical_url("https://www.x.example/a/?utm_source=openai&b=1#f"),
            "https://x.example/a?b=1")
        self.assertEqual(canonical_url("http://X.example/p/"),
                         "https://x.example/p")


class TestIntegrity(_DataTestBase):
    """Regression tests for the pre-publication integrity blockers."""

    def test_incomplete_response_classified_not_ok(self):
        study = self._mk_study(entities=False)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        row = ds.store_observation(man, {
            "observation_id": "obs_inc1", "prompt_id": "S01",
            "run_number": 1, "requested_model": "gpt-5.6-sol",
            "timestamp": "2026-09-01T00:00:00Z", "location": "x",
            "prompt_text": "p", "status": "ok", "error": "",
            "latency_ms": 1}, fx.incomplete_payload(), None, {})
        self.assertEqual(row["status"], "incomplete")
        self.assertIn("max_output_tokens", row["error"])
        ds.update_totals(man)
        man2 = ds.load_manifest(study["study_id"], man["dataset_id"])
        self.assertEqual(man2["totals"]["ok"], 0)
        self.assertEqual(man2["totals"]["failed"], 1)
        checks = exports.reconcile(study["study_id"], man["dataset_id"])
        self.assertTrue(exports.reconciliation_ok(checks),
                        [c for c in checks if not c[1]])

    def test_prompt_location_override(self):
        study = self._mk_study(entities=False)
        study["prompts"][0]["location"] = "Sydney, New South Wales, AU"
        study["locations"] = [{"city": "Melbourne", "region": "Victoria",
                               "country": "AU"},
                              {"city": "Brisbane", "region": "Queensland",
                               "country": "AU"}]
        ds.save_study(study)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        seen = []

        class Client:
            class responses:
                @staticmethod
                def create(**kw):
                    seen.append(kw["tools"][0]["user_location"]["city"])
                    return TestModelIntegrity._FakeResp(fx.plain_payload())
        res = runner.run_batch(Client(), study, man, {})
        self.assertFalse(res["aborted"], res["abort_reason"])
        # Overridden prompt runs ONLY in Sydney; the other prompt runs in
        # both study locations.
        self.assertEqual(seen.count("Sydney"), 1)
        self.assertEqual(seen.count("Melbourne"), 1)
        self.assertEqual(seen.count("Brisbane"), 1)

    def test_mode_enforcement_blocks_run(self):
        study = self._mk_study(entities=True)   # 3 entities
        study["mode"] = "single"
        ds.save_study(study)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        res = runner.run_batch(object(), study, man, {})
        self.assertTrue(res["aborted"])
        self.assertIn("exactly one entity", res["abort_reason"])
        study["mode"] = "unbranded"
        res = runner.run_batch(object(), study, man, {})
        self.assertTrue(res["aborted"])
        self.assertIn("Unbranded", res["abort_reason"])

    def test_duplicate_ids_block_run(self):
        study = self._mk_study(entities=False)
        study["prompts"].append(dict(study["prompts"][0]))  # dup prompt_id
        res = runner.run_batch(object(), study,
                               ds.create_dataset(study, "live",
                                                 "gpt-5.6-sol", 1), {})
        self.assertTrue(res["aborted"])
        self.assertIn("Duplicate prompt ids", res["abort_reason"])

    def test_two_batches_same_dataset_still_reconcile(self):
        # The UI no longer offers appends, but core writes must stay
        # coherent if a dataset receives two batches.
        study = self._mk_study(entities=False)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        client = TestModelIntegrity._FakeClient([fx.plain_payload()])
        runner.run_batch(client, study, man, {})
        runner.run_batch(client, study, man, {})
        man2 = ds.load_manifest(study["study_id"], man["dataset_id"])
        self.assertEqual(man2["totals"]["observations"], 4)  # 2 prompts x2
        checks = exports.reconcile(study["study_id"], man["dataset_id"])
        self.assertTrue(exports.reconciliation_ok(checks),
                        [c for c in checks if not c[1]])

    def test_response_id_stored_and_checked(self):
        study = self._mk_study()
        man = self._fill(study)
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        rids = [r[0] for r in conn.execute(
            "SELECT response_id FROM observations WHERE status='ok'")]
        conn.close()
        self.assertTrue(all(rids), rids)
        checks = dict((n, ok) for n, ok, _ in exports.reconcile(
            study["study_id"], man["dataset_id"]))
        self.assertTrue(checks["ok rows carry response_id"])
        self.assertTrue(checks["frozen prompt ids unique"])
        self.assertTrue(checks["frozen entity ids unique"])

    def test_full_dumps_build(self):
        # Regression: the MD dump crashed with KeyError 'first_party_source'
        # after the per-stage rename; dumps must always be exercised.
        study = self._mk_study()
        man = self._fill(study)
        md = exports.full_markdown_bytes(study["study_id"],
                                         man["dataset_id"]).decode("utf-8")
        for probe in ("source returned", "page opened", "page searched",
                      "named in answer", "Retrieval trail",
                      "No reasoning summary returned"):
            self.assertIn(probe, md)
        self.assertNotIn("first_party_source", md)
        import json as _j
        doc = _j.loads(exports.full_json_bytes(study["study_id"],
                                               man["dataset_id"]))
        self.assertIn("source_entity", doc["tables"])
        self.assertTrue(all(c["ok"] for c in doc["reconciliation"]),
                        doc["reconciliation"])

    def test_source_entity_evidence_stage(self):
        ex = extract_observation(fx.rich_payload(), _matcher())
        stages = {(r["entity_id"], r["evidence_stage"])
                  for r in ex["source_entity"] if r["relation"] == "first_party"}
        self.assertIn(("ent_acme", "search_source"), stages)
        self.assertIn(("ent_acme", "opened_page"), stages)
        self.assertIn(("ent_acme", "citation"), stages)
        self.assertIn(("ent_rival", "inpage_search"), stages)
        # Rival was never opened via open_page:
        self.assertNotIn(("ent_rival", "opened_page"), stages)

    def test_answer_mention_rate_naming(self):
        study = self._mk_study()
        man = self._fill(study)
        conn = open_db(ds.dataset_db_path(study["study_id"],
                                          man["dataset_id"]))
        ev = analysis.entity_visibility(conn, study["entities"])
        conn.close()
        self.assertIn("answer_mention_rate", ev[0])
        self.assertIn("source_returned_rate", ev[0])
        self.assertIn("page_searched_rate", ev[0])
        self.assertNotIn("recommendation_rate", ev[0])


class TestNewCapture(_DataTestBase):
    """Round-3 additions: citation occurrences, search-call provenance,
    response provenance, failed bodies, advanced settings, insights report."""

    def test_citation_occurrences(self):
        ex = extract_observation(fx.rich_payload(), _matcher())
        occ = ex["citation_occurrences"]
        self.assertEqual(len(occ), 3)                       # every occurrence
        self.assertEqual(sum(c["citation_count"]
                             for c in ex["citations"]), 3)  # parity
        first = occ[0]
        self.assertEqual(first["occurrence_number"], 1)
        self.assertIsNotNone(first["start_index"])
        self.assertTrue(first["answer_sentence"])
        self.assertEqual(first["cited_text"],
                         "Acme SEO stands out locally; Rival Digital is "
                         "also frequently recommended."[
                             first["start_index"]:first["end_index"]])

    def test_search_call_provenance(self):
        ex = extract_observation(fx.rich_payload(), _matcher())
        calls = ex["search_calls"]
        self.assertEqual(len(calls), 5)  # 3 search + open + find
        self.assertTrue(all(c["search_call_id"].startswith("ws_fixture_")
                            for c in calls))
        self.assertTrue(all(c["search_call_status"] == "completed"
                            for c in calls))
        srch = [c for c in calls if c["action_type"] == "search"]
        self.assertEqual([c["query_count"] for c in srch], [3, 1, 0])
        self.assertEqual(srch[0]["source_count"], 4)

    def test_failed_body_keeps_api_error(self):
        study = self._mk_study(entities=False)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        row = ds.store_observation(man, {
            "observation_id": "obs_fail1", "prompt_id": "S01",
            "run_number": 1, "requested_model": "gpt-5.6-sol",
            "timestamp": "2026-09-03T00:00:00Z", "location": "x",
            "prompt_text": "p", "status": "ok", "error": "",
            "latency_ms": 1}, fx.failed_payload(), None, {})
        self.assertEqual(row["status"], "error")
        self.assertIn("server_error", row["error"])
        self.assertIn("fixture", row["error"])
        self.assertEqual(row["api_error_code"], "server_error")

    def test_advanced_settings_reach_the_api(self):
        study = self._mk_study(entities=False)
        man = ds.create_dataset(
            study, "live", "gpt-5.6-sol", 1,
            search_settings={"search_context_size": "high",
                             "tool_choice": "required",
                             "max_tool_calls": 7,
                             "allowed_domains": ["example.com"]})
        seen = {}

        class Client:
            class responses:
                @staticmethod
                def create(**kw):
                    seen.update(kw)
                    return TestModelIntegrity._FakeResp(fx.plain_payload())
        res = runner.run_batch(Client(), study, man, {})
        self.assertFalse(res["aborted"], res["abort_reason"])
        self.assertEqual(seen["tools"][0]["search_context_size"], "high")
        self.assertEqual(seen["tools"][0]["filters"]["allowed_domains"],
                         ["example.com"])
        self.assertEqual(seen["tool_choice"], "required")
        self.assertEqual(seen["max_tool_calls"], 7)
        man2 = ds.load_manifest(study["study_id"], man["dataset_id"])
        self.assertEqual(man2["search_settings"]["search_context_size"],
                         "high")

    def test_insights_report_single_file(self):
        study = self._mk_study()
        man = self._fill(study)
        md = exports.insights_markdown_bytes(
            study["study_id"], man["dataset_id"]).decode("utf-8")
        for probe in ("insights", "Key numbers", "Entity visibility",
                      "Prompt coverage", "Query clusters",
                      "Domain leaderboard"):
            self.assertIn(probe, md)
        self.assertNotIn("Retrieval trail (chronological)", md)  # no traces


class TestPublicHardening(_DataTestBase):
    def test_export_formula_injection_escaped(self):
        study = self._mk_study(entities=False)
        study["prompts"] = [__import__("core.prompts", fromlist=["new_prompt"])
                            .new_prompt("=SUM(A1:A9)+cmd", prompt_id="P01")]
        ds.save_study(study)
        man = ds.create_dataset(study, "live", "gpt-5.6-sol", 1)
        # store one obs whose prompt_text is dangerous
        ds.store_observation(man, {
            "observation_id": "obs_x", "prompt_id": "P01", "run_number": 1,
            "requested_model": "gpt-5.6-sol", "timestamp": "t",
            "location": "x", "prompt_text": "=HYPERLINK(\"evil\")",
            "status": "ok", "error": "", "latency_ms": 1},
            fx.plain_payload(), None, {})
        csv_bytes = exports.table_csv_bytes(study["study_id"],
                                            man["dataset_id"], "observations")
        text = csv_bytes.decode("utf-8-sig")
        # the dangerous cell is prefixed with a quote, not left executable
        self.assertNotIn(",=HYPERLINK", text)
        self.assertIn("'=HYPERLINK", text)

    def test_upload_rejects_macro_and_oversize(self):
        from core.prompts import parse_upload, UploadError, MAX_UPLOAD_BYTES
        import io as _io

        class F(_io.BytesIO):
            def __init__(self, data, name):
                super().__init__(data); self.name = name
        with self.assertRaises(UploadError):
            parse_upload(F(b"x", "prompts.xlsm"))
        with self.assertRaises(UploadError):
            parse_upload(F(b"x" * (MAX_UPLOAD_BYTES + 1), "big.csv"))
        with self.assertRaises(UploadError):
            parse_upload(F(b"not a real xlsx", "bad.xlsx"))  # corrupt

    def test_upload_caps_prompt_count(self):
        from core.prompts import parse_upload, MAX_IMPORT_PROMPTS
        import io as _io

        class F(_io.BytesIO):
            name = "many.txt"
        big = "\n".join(f"prompt {i}"
                        for i in range(MAX_IMPORT_PROMPTS + 50))
        ps, _u, _b = parse_upload(F(big.encode()))
        self.assertLessEqual(len(ps), MAX_IMPORT_PROMPTS)


class TestHostingIsolation(unittest.TestCase):
    """Per-session isolation foundation for the hosted deploy mode."""

    def setUp(self):
        import core.hosting as hosting
        self.hosting = hosting
        self._tmp = tempfile.mkdtemp(prefix="fanout_host_")
        self._old = hosting.SESSIONS_ROOT
        hosting.SESSIONS_ROOT = os.path.join(self._tmp, "sessions")

    def tearDown(self):
        ds.set_session_root(None)
        self.hosting.SESSIONS_ROOT = self._old
        shutil.rmtree(self._tmp, ignore_errors=True)

    def _make_study_in_session(self, sid, name):
        ds.set_session_root(self.hosting.session_workspace(sid))
        st = ds.create_study(name, mode="unbranded")
        return st

    def test_two_sessions_cannot_see_each_other(self):
        self._make_study_in_session("sessAAA", "Alice study")
        self._make_study_in_session("sessBBB", "Bob study")
        # Alice's view
        ds.set_session_root(self.hosting.session_workspace("sessAAA"))
        names_a = [x["study_name"] for x in ds.list_studies()]
        # Bob's view
        ds.set_session_root(self.hosting.session_workspace("sessBBB"))
        names_b = [x["study_name"] for x in ds.list_studies()]
        self.assertEqual(names_a, ["Alice study"])
        self.assertEqual(names_b, ["Bob study"])
        self.assertNotIn("Alice study", names_b)

    def test_clear_session_removes_files(self):
        self._make_study_in_session("sessX", "Temp study")
        ws = self.hosting.session_workspace("sessX")
        self.assertTrue(os.listdir(ws))                 # study written
        self.assertTrue(self.hosting.clear_session("sessX"))
        self.assertFalse(os.path.exists(
            os.path.dirname(ws)))                        # whole dir gone

    def test_gc_removes_expired_only(self):
        self._make_study_in_session("old", "Old")
        self._make_study_in_session("new", "New")
        old_dir = os.path.join(self.hosting.SESSIONS_ROOT, "old")
        past = time.time() - 48 * 3600
        os.utime(old_dir, (past, past))
        removed = self.hosting.gc_sessions(ttl_hours=12)
        self.assertEqual(removed, 1)
        self.assertFalse(os.path.exists(old_dir))
        self.assertTrue(os.path.exists(
            os.path.join(self.hosting.SESSIONS_ROOT, "new")))

    def test_local_mode_unaffected(self):
        # With no session root set, storage falls back to DATA_ROOT (local).
        ds.set_session_root(None)
        self.assertEqual(ds.data_root(), ds.DATA_ROOT)

    def test_sid_path_traversal_is_neutralised(self):
        ws = self.hosting.session_workspace("../../etc/evil")
        self.assertTrue(os.path.abspath(ws).startswith(
            os.path.abspath(self.hosting.SESSIONS_ROOT)))


class TestLeadCapture(unittest.TestCase):
    def setUp(self):
        import core.leads as leads
        self.leads = leads
        self._tmp = tempfile.mkdtemp(prefix="fanout_leads_")
        self._old = leads.LEADS_PATH
        leads.LEADS_PATH = os.path.join(self._tmp, "leads.jsonl")

    def tearDown(self):
        self.leads.LEADS_PATH = self._old
        shutil.rmtree(self._tmp, ignore_errors=True)

    def test_email_validation(self):
        self.assertTrue(self.leads.valid_email("a@b.co"))
        self.assertFalse(self.leads.valid_email("nope"))
        self.assertFalse(self.leads.valid_email("a@b"))

    def test_record_lead_appends_and_never_stores_key(self):
        self.leads.record_lead("Ada", "ada@example.com", "Acme",
                               session_id="s1",
                               extra={"api_key": "sk-secret", "note": "hi"})
        rows = self.leads.read_leads()
        self.assertEqual(len(rows), 1)
        r = rows[0]
        self.assertEqual(r["name"], "Ada")
        self.assertEqual(r["email"], "ada@example.com")
        self.assertEqual(r["note"], "hi")
        # forbidden fields stripped; no sk- anywhere in the record
        self.assertNotIn("api_key", r)
        self.assertNotIn("sk-secret", json.dumps(r))

    def test_webhook_gets_token_but_local_record_does_not(self):
        import os as _os, urllib.request as _u, json as _json
        captured = {}

        class _Resp:
            def read(self_):  # noqa: N805
                return b"ok"

        def fake_urlopen(req, timeout=0):
            captured["body"] = _json.loads(req.data.decode("utf-8"))
            return _Resp()

        old_open = _u.urlopen
        _os.environ["RETRIEVAL_LEADS_WEBHOOK"] = "https://example.test/hook"
        _os.environ["RETRIEVAL_LEADS_TOKEN"] = "s3cret"
        _u.urlopen = fake_urlopen
        try:
            rec = self.leads.record_lead("Ada", "ada@example.com")
        finally:
            _u.urlopen = old_open
            _os.environ.pop("RETRIEVAL_LEADS_WEBHOOK", None)
            _os.environ.pop("RETRIEVAL_LEADS_TOKEN", None)
        # token reached the webhook payload...
        self.assertEqual(captured["body"].get("_token"), "s3cret")
        self.assertEqual(captured["body"].get("email"), "ada@example.com")
        # ...but is NOT in the returned record or the local jsonl
        self.assertNotIn("_token", rec)
        self.assertNotIn("s3cret",
                         _json.dumps(self.leads.read_leads()))

    def test_multiple_leads_accumulate(self):
        self.leads.record_lead("A", "a@x.com")
        self.leads.record_lead("B", "b@x.com")
        self.assertEqual(len(self.leads.read_leads()), 2)


class TestModelPricingCoverage(unittest.TestCase):
    """Every selectable model must have a priced, valid id.

    The runner aborts a batch if the requested model id is not echoed back,
    so an unpriced or misspelled id (e.g. the non-existent 'gpt-5.6-astra')
    would silently break every run. Parse the MODELS literal statically to
    avoid importing app.py (which runs Streamlit at module level).
    """

    @classmethod
    def setUpClass(cls):
        import ast
        import json
        here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        src = open(os.path.join(here, "app.py"), encoding="utf-8").read()
        cls.models = None
        for node in ast.walk(ast.parse(src)):
            if (isinstance(node, ast.Assign)
                    and any(getattr(t, "id", None) == "MODELS"
                            for t in node.targets)):
                cls.models = ast.literal_eval(node.value)
                break
        with open(os.path.join(here, "pricing.json"), encoding="utf-8") as fh:
            cls.pricing = json.load(fh)

    def test_models_extracted(self):
        self.assertTrue(self.models, "MODELS list not found in app.py")

    def test_every_model_priced(self):
        for m in self.models:
            with self.subTest(model=m):
                rate = self.pricing.get(m)
                self.assertIsNotNone(rate, f"{m} missing from pricing.json")
                self.assertGreater(rate["input_per_mtok"], 0)
                self.assertGreater(rate["output_per_mtok"], 0)

    def test_no_fictional_astra_id(self):
        # Astra is gpt-6-astra (its own family), never a gpt-5.6 variant.
        self.assertNotIn("gpt-5.6-astra", self.models)
        self.assertNotIn("gpt-5.6-astra", self.pricing)


class TestRepoHygiene(unittest.TestCase):
    """No live research data or secrets in the tracked tree."""

    @classmethod
    def setUpClass(cls):
        import subprocess
        here = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        try:
            out = subprocess.check_output(["git", "ls-files"], cwd=here,
                                          text=True)
            cls.tracked = out.splitlines()
        except Exception:
            cls.tracked = None
        cls.here = here

    def test_no_data_artifacts_tracked(self):
        if self.tracked is None:
            self.skipTest("not a git checkout")
        bad = [f for f in self.tracked
               if f.startswith(("data/", "output/", "sample_output/"))
               or f.endswith((".db", ".log"))
               or "/raw/" in f]
        self.assertEqual(bad, [])

    def test_no_api_keys_tracked(self):
        if self.tracked is None:
            self.skipTest("not a git checkout")
        import re
        pat = re.compile(r"sk-[A-Za-z0-9_\-]{20,}")
        hits = []
        for f in self.tracked:
            p = os.path.join(self.here, f)
            if not os.path.exists(p) or os.path.getsize(p) > 2_000_000:
                continue
            try:
                text = open(p, encoding="utf-8", errors="ignore").read()
            except Exception:
                continue
            if pat.search(text):
                hits.append(f)
        self.assertEqual(hits, [])


if __name__ == "__main__":
    unittest.main(verbosity=2)
